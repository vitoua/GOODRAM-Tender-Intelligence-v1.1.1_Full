from datetime import datetime,timedelta
from io import BytesIO
import asyncio,json
from fastapi import FastAPI,Form,Request,HTTPException
from fastapi.responses import RedirectResponse,StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from app.config import settings
from app.core import store,active,analytics
from app.sources import COUNTRIES
from app.i18n import tr
from app.translate import ensure
from pathlib import Path
app=FastAPI(version='1.1.1');app.mount('/static',StaticFiles(directory='app/static'),name='static');tpl=Jinja2Templates(directory='app/templates')
CPV=Path('data/cpv.json');KW=Path('data/keywords.json')
def lang(r):return r.cookies.get('lang') or settings.app_language
def ctx(r,x=None):return {'request':r,'lang':lang(r),'t':tr(lang(r)),**(x or {})}
def default_date():return (datetime.utcnow()-timedelta(days=30)).date().isoformat()
@app.get('/health')
def health():return {'status':'ok','version':'1.1.1','tenders':len(store.data)}
@app.get('/lang/{code}')
def language(code,request:Request):
 r=RedirectResponse(request.headers.get('referer') or '/',303);r.set_cookie('lang',code if code in ('pl','uk','en','es') else 'pl');return r
@app.get('/')
async def home(request:Request,q:str='',country:str='',source:str='',active_only:int=1):
 rows=[x for x in store.data.values() if (not active_only or active(x)) and (not country or x.country==country) and (not source or x.source==source) and (not q or q.casefold() in (x.title+x.buyer+x.description+' '.join(x.cpv)).casefold())];rows.sort(key=lambda x:x.deadline or datetime.max);await ensure(rows,lang(request));return tpl.TemplateResponse('index.html',ctx(request,{'rows':rows,'runs':store.runs,'countries':COUNTRIES,'default_date':default_date(),'q':q,'country':country,'source':source,'active_only':active_only}))
@app.post('/refresh')
async def refresh(countries:list[str]=Form(default=[]),min_value_eur:float=Form(50000),published_from:str=Form(default='')):
 selected=countries or list(COUNTRIES);start=datetime.fromisoformat(published_from or default_date());tasks=[];eu=[c for c in selected if c!='UKR']
 if eu:
  from app.connectors.ted import fetch;tasks.append(('ted',fetch(eu,start,min_value_eur)))
 if 'UKR' in selected:
  from app.connectors.prozorro import fetch;tasks.append(('prozorro',fetch(['UKR'],start,min_value_eur)))
 if 'DEU' in selected:
  from app.connectors.germany import fetch;tasks.append(('germany',fetch(['DEU'],start,min_value_eur)))
 async def run(n,c):
  try:return n,await asyncio.wait_for(c,100),None
  except Exception as e:return n,[],e
 for n,rows,e in await asyncio.gather(*(run(n,c) for n,c in tasks)):store.error(n,e) if e else store.add(n,rows)
 return RedirectResponse('/',303)
@app.get('/tenders/{tid:path}')
async def detail(tid,request:Request,view:str='translated'):
 x=store.data.get(tid)
 if not x:raise HTTPException(404)
 await ensure([x],lang(request));return tpl.TemplateResponse('detail.html',ctx(request,{'x':x,'view':view}))
@app.get('/analytics')
def analytics_page(request:Request,country:str='',source:str='',date_from:str='',date_to:str=''):
 rows=list(store.data.values())
 if country:rows=[x for x in rows if x.country==country]
 if source:rows=[x for x in rows if x.source==source]
 if date_from:rows=[x for x in rows if (x.published or x.deadline) and (x.published or x.deadline).date()>=datetime.fromisoformat(date_from).date()]
 if date_to:rows=[x for x in rows if (x.published or x.deadline) and (x.published or x.deadline).date()<=datetime.fromisoformat(date_to).date()]
 return tpl.TemplateResponse('analytics.html',ctx(request,{'data':analytics(rows),'count':len(rows),'country':country,'source':source,'date_from':date_from,'date_to':date_to,'countries':sorted({x.country for x in store.data.values()}),'source_list':sorted({x.source for x in store.data.values()})}))
@app.get('/cpv')
def cpv_page(request:Request):return tpl.TemplateResponse('cpv.html',ctx(request,{'rows':json.loads(CPV.read_text())}))
@app.post('/cpv/save')
def cpv_save(code:list[str]=Form(default=[]),name:list[str]=Form(default=[]),active_codes:list[str]=Form(default=[])):
 enabled=set(active_codes);CPV.write_text(json.dumps([{'code':c.strip(),'name_pl':n.strip(),'active':c in enabled} for c,n in zip(code,name) if c.strip()],ensure_ascii=False,indent=2));return RedirectResponse('/cpv',303)
@app.post('/cpv/add')
def cpv_add(code:str=Form(),name:str=Form()):
 rows=json.loads(CPV.read_text());rows.append({'code':code.strip(),'name_pl':name.strip(),'active':True});CPV.write_text(json.dumps(rows,ensure_ascii=False,indent=2));return RedirectResponse('/cpv',303)
@app.get('/keywords')
def keyword_page(request:Request):return tpl.TemplateResponse('keywords.html',ctx(request,{'groups':json.loads(KW.read_text())}))
@app.post('/keywords/save')
def keywords_save(pl:str=Form(''),uk:str=Form(''),en:str=Form('')):
 data={k:[x.strip() for x in v.splitlines() if x.strip()] for k,v in {'pl':pl,'uk':uk,'en':en}.items()};KW.write_text(json.dumps(data,ensure_ascii=False,indent=2));return RedirectResponse('/keywords',303)
@app.get('/sources')
def sources_page(request:Request):return tpl.TemplateResponse('sources.html',ctx(request,{'countries':COUNTRIES}))
@app.get('/about')
def about(request:Request):return tpl.TemplateResponse('about.html',ctx(request))
@app.get('/export.xlsx')
def export_excel(q:str='',country:str='',source:str=''):
 from openpyxl import Workbook
 from openpyxl.styles import Font,PatternFill,Alignment
 from openpyxl.utils import get_column_letter
 from openpyxl.worksheet.table import Table,TableStyleInfo
 rows=[x for x in store.data.values() if (not country or x.country==country) and (not source or x.source==source) and (not q or q.casefold() in (x.title+x.buyer+x.description+' '.join(x.cpv)).casefold())]
 wb=Workbook();ws=wb.active;ws.title='Przetargi';headers=['Źródło','Kraj','Numer postępowania','Nazwa','Zamawiający','Data publikacji','Termin składania ofert','Wartość','Waluta','Wartość EUR','CPV','Status','Link'];ws.append(headers)
 for x in sorted(rows,key=lambda x:x.deadline or datetime.max):
  ws.append([x.source.upper(),x.country,x.procurement_id,x.title,x.buyer,x.published,x.deadline,x.value,x.currency,x.value_eur,', '.join(x.cpv),x.status,'Otwórz przetarg']);cell=ws.cell(ws.max_row,13);cell.hyperlink=x.url;cell.style='Hyperlink'
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='E30613');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(vertical='center')
 ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
 if ws.max_row>1:
  tab=Table(displayName='TabelaPrzetargow',ref=f'A1:M{ws.max_row}');tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab)
 for i,w in enumerate([14,10,22,55,38,18,22,16,10,16,34,14,18],1):ws.column_dimensions[get_column_letter(i)].width=w
 for r in range(2,ws.max_row+1):ws.cell(r,6).number_format='yyyy-mm-dd';ws.cell(r,7).number_format='yyyy-mm-dd hh:mm';ws.cell(r,8).number_format='#,##0.00';ws.cell(r,10).number_format='#,##0.00'
 out=BytesIO();wb.save(out);out.seek(0);return StreamingResponse(out,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="GOODRAM-Tender-Intelligence-export.xlsx"'})
